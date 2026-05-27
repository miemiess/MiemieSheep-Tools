"""
高速转发功能模块 —— MQTT 客户端
支持自定义消息发送 + 跨平台 OTA 升级（单设备 / Excel 批量）
"""
import datetime
import json
import os

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QLineEdit, QTextEdit, QGroupBox,
    QFormLayout, QSpinBox, QFrame, QSizePolicy,
    QFileDialog, QStackedWidget, QScrollArea,
)
from PySide6.QtCore import Qt, QObject, Signal, QTimer

CONN_TIMEOUT_MS = 3000  # 连接超时：3 秒


# ─── OTA payload 构建 ─────────────────────────────────────────────────────────
def _build_ota_payload(url: str, ver: str, size: int, sha256: str,
                       pretty: bool = False) -> str:
    """构建 OTA JSON payload（与原 TEMPLATE_1_PAYLOAD 结构相同）"""
    return json.dumps(
        {
            "cmd": "ota",
            "requestId": "ota_001",
            "data": {
                "espfileUrl":    url,
                "espver":        ver,
                "espfilesize":   size,
                "espfilesha256": sha256,
            },
        },
        ensure_ascii=False,
        indent=2 if pretty else None,
    )


# ─── 按钮辅助 ─────────────────────────────────────────────────────────────────
_BTN_CSS = """
    QPushButton {{
        background-color: {c}; color: white; border: none;
        border-radius: 7px; padding: 3px 14px;
        font-size: 12px; font-weight: bold;
    }}
    QPushButton:hover    {{ background-color: {ch}; }}
    QPushButton:disabled {{ background-color: {cd}; color: #bbb; }}
"""


def _make_btn(text: str, hex_color: str, height: int = 30) -> QPushButton:
    btn = QPushButton(text)
    btn.setFixedHeight(height)
    btn.setCursor(Qt.CursorShape.PointingHandCursor)
    btn.setStyleSheet(_BTN_CSS.format(
        c=hex_color, ch=hex_color + "CC", cd=hex_color + "55"))
    return btn


def _outline_btn(text: str, color: str, height: int = 24) -> QPushButton:
    btn = QPushButton(text)
    btn.setFixedHeight(height)
    btn.setCursor(Qt.CursorShape.PointingHandCursor)
    btn.setStyleSheet(f"""
        QPushButton {{
            background-color: transparent; color: {color};
            border: 1px solid {color}; border-radius: 5px;
            padding: 1px 8px; font-size: 11px;
        }}
        QPushButton:hover {{ background-color: {color}22; }}
    """)
    return btn


# ─── MQTTManager ──────────────────────────────────────────────────────────────
class MQTTManager(QObject):
    sig_connected    = Signal()
    sig_disconnected = Signal()
    sig_log          = Signal(str)   # HTML 片段
    sig_publish_ok   = Signal()      # publish rc == 0
    sig_publish_fail = Signal()      # publish failed

    def __init__(self, parent=None):
        super().__init__(parent)
        self._client = None
        self._conn_timer = QTimer(self)
        self._conn_timer.setSingleShot(True)
        self._conn_timer.setInterval(CONN_TIMEOUT_MS)
        self._conn_timer.timeout.connect(self._on_conn_timeout)

    # ---- 连接 --------------------------------------------------------
    def connect_broker(self, host: str, port: int, client_id: str,
                       username: str, password: str, keepalive: int):
        self.disconnect_broker()
        try:
            import paho.mqtt.client as mqtt
        except ImportError:
            self.sig_log.emit(
                self._fmt("[ERROR]", "paho-mqtt 未安装，请执行: pip install paho-mqtt", "red")
            )
            self.sig_disconnected.emit()
            return
        try:
            try:
                client = mqtt.Client(
                    mqtt.CallbackAPIVersion.VERSION1,
                    client_id=client_id,
                    clean_session=True,
                )
            except AttributeError:
                client = mqtt.Client(client_id=client_id, clean_session=True)

            if username:
                client.username_pw_set(username, password or "")

            client.on_connect    = self._on_connect
            client.on_disconnect = self._on_disconnect
            client.on_message    = self._on_message

            self._client = client
            self._client.connect_async(host, port, keepalive)
            self._client.loop_start()
            self._conn_timer.start()
            self.sig_log.emit(
                self._fmt("[INFO]",
                          f"正在连接 {host}:{port}（{CONN_TIMEOUT_MS // 1000} 秒超时）...",
                          "cyan")
            )
        except Exception as exc:
            self.sig_log.emit(self._fmt("[ERROR]", f"连接异常: {exc}", "red"))
            self.sig_disconnected.emit()

    def _on_conn_timeout(self):
        self.sig_log.emit(
            self._fmt("[ERROR]",
                      f"连接超时（{CONN_TIMEOUT_MS // 1000} 秒），已自动断开，请检查地址与网络",
                      "red")
        )
        self.disconnect_broker()
        self.sig_disconnected.emit()

    # ---- 断开 --------------------------------------------------------
    def disconnect_broker(self):
        self._conn_timer.stop()
        if self._client is not None:
            try:
                self._client.loop_stop()
                self._client.disconnect()
            except Exception:
                pass
            self._client = None

    # ---- 发布 --------------------------------------------------------
    def publish(self, topic: str, payload: str):
        if self._client is None:
            self.sig_log.emit(self._fmt("[WARN]", "尚未连接到 MQTT 服务器", "orange"))
            self.sig_publish_fail.emit()
            return
        try:
            result = self._client.publish(topic, payload)
            if result.rc == 0:
                short = payload if len(payload) <= 120 else payload[:120] + "…"
                self.sig_log.emit(
                    self._fmt("[SENT ✓]",
                              f"Topic: {topic}<br>&nbsp;&nbsp;Payload: {short}",
                              "#5EF09E")
                )
                self.sig_publish_ok.emit()
            else:
                self.sig_log.emit(
                    self._fmt("[ERROR]", f"发送失败，rc={result.rc}", "red")
                )
                self.sig_publish_fail.emit()
        except Exception as exc:
            self.sig_log.emit(self._fmt("[ERROR]", f"发送异常: {exc}", "red"))
            self.sig_publish_fail.emit()

    def is_connected(self) -> bool:
        return self._client is not None and self._client.is_connected()

    # ---- paho 回调（在 paho 内部线程，Qt Signal 自动跨线程安全）-------
    def _on_connect(self, client, userdata, flags, rc):
        self._conn_timer.stop()
        _CODES = {1: "协议版本错误", 2: "Client ID 无效",
                  3: "服务器不可用", 4: "用户名/密码错误", 5: "未授权"}
        if rc == 0:
            self.sig_connected.emit()
            self.sig_log.emit(self._fmt("[CONN ✓]", "成功连接到 MQTT 服务器", "#5EF09E"))
        else:
            msg = _CODES.get(rc, f"未知错误码 {rc}")
            self.sig_log.emit(self._fmt("[ERROR]", f"连接失败: {msg}", "red"))
            self.sig_disconnected.emit()

    def _on_disconnect(self, client, userdata, rc):
        self._conn_timer.stop()
        self.sig_disconnected.emit()
        if rc == 0:
            self.sig_log.emit(self._fmt("[INFO]", "已主动断开连接", "cyan"))
        else:
            self.sig_log.emit(
                self._fmt("[WARN]", f"意外断开，rc={rc}，请检查网络", "orange")
            )

    def _on_message(self, client, userdata, msg):
        try:
            payload = msg.payload.decode("utf-8")
        except Exception:
            payload = repr(msg.payload)
        self.sig_log.emit(
            self._fmt("[RECV]",
                      f"Topic: {msg.topic}<br>&nbsp;&nbsp;Payload: {payload}",
                      "#56CCF2")
        )

    @staticmethod
    def _fmt(tag: str, body: str, color: str) -> str:
        ts = datetime.datetime.now().strftime("%H:%M:%S")
        return (
            f'<span style="color:#888;">[{ts}]</span> '
            f'<span style="color:{color};font-weight:bold;">{tag}</span> '
            f'<span style="color:#C8E6D5;">{body}</span>'
        )


# ─── ExcelBatchManager ────────────────────────────────────────────────────────
class ExcelBatchManager:
    """读取 Excel 中 uid 列，发送成功后将对应单元格标绿保存"""

    def __init__(self):
        self._path: str | None = None
        self._wb = None
        self._ws = None
        self._uid_col: int | None = None        # 1-based 列号
        self._uid_rows: list[tuple[int, str]] = []  # (行号, uid值)

    def import_file(self, path: str) -> tuple[bool, str]:
        try:
            import openpyxl
        except ImportError:
            return False, "请先安装 openpyxl：pip install openpyxl"
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            # 在第一行找 'uid' 列（不区分大小写）
            uid_col = None
            for cell in ws[1]:
                if cell.value and str(cell.value).strip().lower() == "uid":
                    uid_col = cell.column
                    break
            if uid_col is None:
                wb.close()
                return False, "第一行中未找到 'uid' 列标题，请检查 Excel 文件"
            # 从第 2 行起读取所有非空 UID
            uid_rows: list[tuple[int, str]] = []
            for row in range(2, ws.max_row + 1):
                val = ws.cell(row=row, column=uid_col).value
                if val is not None and str(val).strip():
                    uid_rows.append((row, str(val).strip()))
            if not uid_rows:
                wb.close()
                return False, "uid 列中无有效数据（第 2 行起）"
            self.reset()
            self._path    = path
            self._wb      = wb
            self._ws      = ws
            self._uid_col = uid_col
            self._uid_rows = uid_rows
            return True, f"成功读取 {len(uid_rows)} 个 UID"
        except Exception as exc:
            return False, f"读取 Excel 失败: {exc}"

    def mark_green(self, list_idx: int):
        """将第 list_idx 个 UID 单元格标绿并立即保存"""
        if self._wb is None or list_idx >= len(self._uid_rows):
            return
        try:
            from openpyxl.styles import PatternFill
            row_idx, _ = self._uid_rows[list_idx]
            cell = self._ws.cell(row=row_idx, column=self._uid_col)
            cell.fill = PatternFill(fill_type="solid", fgColor="92D050")
            self._wb.save(self._path)
        except Exception:
            pass  # 标色失败不中断主流程

    def get_uid(self, idx: int) -> str:
        return self._uid_rows[idx][1] if 0 <= idx < len(self._uid_rows) else ""

    @property
    def count(self) -> int:
        return len(self._uid_rows)

    @property
    def file_name(self) -> str:
        return os.path.basename(self._path) if self._path else ""

    def reset(self):
        if self._wb:
            try:
                self._wb.close()
            except Exception:
                pass
        self._path     = None
        self._wb       = None
        self._ws       = None
        self._uid_col  = None
        self._uid_rows = []


# ─── BatchExecutionPanel ──────────────────────────────────────────────────────
class BatchExecutionPanel(QFrame):
    """
    Excel 批量 UID → OTA payload 发送面板
    状态机: empty → imported → sending → sent_ok | sent_fail → done
    get_payload_fn: callable() -> str | None  （返回已确认的 OTA payload）
    """

    def __init__(self, mqtt_mgr: "MQTTManager",
                 get_payload_fn,
                 parent=None):
        super().__init__(parent)
        self._mqtt         = mqtt_mgr
        self._get_payload  = get_payload_fn
        self._excel        = ExcelBatchManager()
        self._idx          = 0
        self._batch_active = False

        self._mqtt.sig_publish_ok.connect(self._on_publish_ok)
        self._mqtt.sig_publish_fail.connect(self._on_publish_fail)
        self._mqtt.sig_connected.connect(self._refresh_start_btn)
        self._mqtt.sig_disconnected.connect(self._refresh_start_btn)

        self._setup_ui()
        self._set_state("empty")

    # ── UI ────────────────────────────────────────────────────────────
    def _setup_ui(self):
        self.setObjectName("batchPanelFrame")
        self.setStyleSheet("""
            QFrame#batchPanelFrame {
                background-color: #F5FFFB;
                border-radius: 10px;
                border: 2px solid #C8E6D5;
            }
        """)
        outer = QVBoxLayout(self)
        outer.setContentsMargins(12, 8, 12, 10)
        outer.setSpacing(5)

        # 标题行
        hdr = QHBoxLayout()
        title_lbl = QLabel("📊  Excel 批量 OTA 推送")
        title_lbl.setStyleSheet("font-weight:bold; font-size:12px; color:#1A6B45;")
        self._reimport_btn = _outline_btn("重新导入", "#5B9BD5")
        self._reimport_btn.clicked.connect(self._do_import)
        self._reimport_btn.hide()
        hdr.addWidget(title_lbl)
        hdr.addStretch()
        hdr.addWidget(self._reimport_btn)
        outer.addLayout(hdr)

        # 文件 + 进度行
        info_row = QHBoxLayout()
        self._file_lbl = QLabel("未导入文件")
        self._file_lbl.setStyleSheet("font-size:11px; color:#8BAF9B; font-family:Consolas;")
        self._progress_lbl = QLabel("")
        self._progress_lbl.setStyleSheet("font-size:11px; color:#5B9BD5; font-weight:bold;")
        info_row.addWidget(self._file_lbl)
        info_row.addStretch()
        info_row.addWidget(self._progress_lbl)
        outer.addLayout(info_row)

        # 当前 UID + 状态行
        uid_row = QHBoxLayout()
        uid_prefix = QLabel("当前 UID:")
        uid_prefix.setStyleSheet("font-size:12px; color:#4A9970;")
        uid_prefix.setFixedWidth(65)
        self._cur_uid_lbl = QLabel("")
        self._cur_uid_lbl.setStyleSheet(
            "font-size:13px; font-weight:bold; color:#1A6B45; font-family:Consolas;")
        self._status_lbl = QLabel("")
        self._status_lbl.setStyleSheet("font-size:12px; color:#3EC78C; font-weight:bold;")
        uid_row.addWidget(uid_prefix)
        uid_row.addWidget(self._cur_uid_lbl, 1)
        uid_row.addWidget(self._status_lbl)
        outer.addLayout(uid_row)

        # 操作按钮行
        btn_row = QHBoxLayout()
        btn_row.setSpacing(7)
        self._import_btn = _make_btn("📂  导入 Excel", "#5B9BD5")
        self._start_btn  = _make_btn("▶  开始执行",   "#3EC78C")
        self._next_btn   = _make_btn("→  下一个",     "#3EC78C")
        self._retry_btn  = _make_btn("↺  重试",       "#FFB347")
        self._skip_btn   = _make_btn("⏭  跳过",       "#9E9E9E")
        self._abort_btn  = _make_btn("■  终止",       "#FF6B6B")
        self._reset_btn  = _make_btn("🔄  重置",       "#5B9BD5")
        self._import_btn.clicked.connect(self._do_import)
        self._start_btn.clicked.connect(self._do_start)
        self._next_btn.clicked.connect(self._do_next)
        self._retry_btn.clicked.connect(self._do_retry)
        self._skip_btn.clicked.connect(self._do_skip)
        self._abort_btn.clicked.connect(self._do_abort)
        self._reset_btn.clicked.connect(self._do_reset)
        for b in (self._import_btn, self._start_btn, self._next_btn,
                  self._retry_btn, self._skip_btn, self._abort_btn, self._reset_btn):
            btn_row.addWidget(b)
        btn_row.addStretch()
        outer.addLayout(btn_row)

    # ── 状态机 ────────────────────────────────────────────────────────
    #               import reimport start  next   retry  skip   abort  reset
    _STATE_VIS = {
        "empty":     (True,  False, False, False, False, False, False, False),
        "imported":  (False, True,  True,  False, False, False, False, False),
        "sending":   (False, False, False, False, False, False, False, False),
        "sent_ok":   (False, True,  False, True,  True,  False, True,  False),
        "sent_fail": (False, True,  False, False, True,  True,  True,  False),
        "done":      (False, True,  False, False, False, False, False, True ),
    }
    _STATE_MSG = {
        "empty":     ("",                         "#3EC78C"),
        "imported":  ("就绪，点击开始执行",          "#3EC78C"),
        "sending":   ("⏳ 发送中...",              "#FFB347"),
        "sent_ok":   ("✅ 发送成功，单元格已标绿",  "#3EC78C"),
        "sent_fail": ("❌ 发送失败",               "#FF6B6B"),
        "done":      ("🎉 全部执行完毕！",           "#3EC78C"),
    }

    def _set_state(self, state: str):
        btns = [self._import_btn, self._reimport_btn, self._start_btn,
                self._next_btn, self._retry_btn, self._skip_btn,
                self._abort_btn, self._reset_btn]
        for btn, vis in zip(btns, self._STATE_VIS.get(state, (False,) * 8)):
            btn.setVisible(vis)
        self._cur_uid_lbl.setVisible(state != "empty")
        text, color = self._STATE_MSG.get(state, ("", "#3EC78C"))
        self._status_lbl.setText(text)
        self._status_lbl.setStyleSheet(f"font-size:12px; color:{color}; font-weight:bold;")

    def _refresh_start_btn(self):
        self._start_btn.setEnabled(self._mqtt.is_connected())

    # ── 事件处理 ──────────────────────────────────────────────────────
    def _do_import(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "选择 Excel 文件", "", "Excel 文件 (*.xlsx *.xls)")
        if not path:
            return
        ok, msg = self._excel.import_file(path)
        if ok:
            self._idx = 0
            self._file_lbl.setText(f"📄 {self._excel.file_name}")
            self._progress_lbl.setText(f"0 / {self._excel.count}")
            self._cur_uid_lbl.setText(self._excel.get_uid(0))
            self._reimport_btn.show()
            self._set_state("imported")
            self._refresh_start_btn()
        else:
            self._status_lbl.setText(f"❌ {msg}")
            self._status_lbl.setStyleSheet("font-size:12px; color:#FF6B6B; font-weight:bold;")

    def _do_start(self):
        if not self._mqtt.is_connected():
            self._status_lbl.setText("⚠️ 请先连接 MQTT 服务器")
            self._status_lbl.setStyleSheet("font-size:12px; color:#FF6B6B; font-weight:bold;")
            return
        if self._get_payload() is None:
            self._status_lbl.setText("⚠️ 请先在上方填写参数并点击「确认 / 预览」")
            self._status_lbl.setStyleSheet("font-size:12px; color:#FF6B6B; font-weight:bold;")
            return
        self._idx = 0
        self._send_current()

    def _send_current(self):
        if self._idx >= self._excel.count:
            self._set_state("done")
            self._progress_lbl.setText(f"{self._excel.count} / {self._excel.count}")
            return
        payload = self._get_payload()
        if payload is None:
            self._status_lbl.setText("⚠️ OTA 参数已失效，请重新确认")
            self._status_lbl.setStyleSheet("font-size:12px; color:#FF6B6B; font-weight:bold;")
            return
        uid   = self._excel.get_uid(self._idx)
        topic = f"/func/4/{uid}/get"
        self._cur_uid_lbl.setText(uid)
        self._progress_lbl.setText(f"{self._idx + 1} / {self._excel.count}")
        self._set_state("sending")
        self._batch_active = True
        self._mqtt.publish(topic, payload)

    def _on_publish_ok(self):
        if not self._batch_active:
            return
        self._batch_active = False
        self._excel.mark_green(self._idx)
        if self._idx >= self._excel.count - 1:
            self._progress_lbl.setText(f"{self._excel.count} / {self._excel.count}")
            self._set_state("done")
        else:
            self._set_state("sent_ok")

    def _on_publish_fail(self):
        if not self._batch_active:
            return
        self._batch_active = False
        self._set_state("sent_fail")

    def _do_next(self):
        self._idx += 1
        self._send_current()

    def _do_retry(self):
        self._send_current()

    def _do_skip(self):
        self._idx += 1
        if self._idx >= self._excel.count:
            self._set_state("done")
            self._progress_lbl.setText(f"{self._excel.count} / {self._excel.count}")
        else:
            self._cur_uid_lbl.setText(self._excel.get_uid(self._idx))
            self._progress_lbl.setText(f"{self._idx + 1} / {self._excel.count}")
            self._set_state("imported")

    def _do_abort(self):
        self._batch_active = False
        self._cur_uid_lbl.setText(self._excel.get_uid(self._idx))
        self._status_lbl.setText("已终止执行")
        self._status_lbl.setStyleSheet("font-size:12px; color:#9E9E9E; font-weight:bold;")
        self._set_state("imported")

    def _do_reset(self):
        self._batch_active = False
        self._excel.reset()
        self._idx = 0
        self._file_lbl.setText("未导入文件")
        self._progress_lbl.setText("")
        self._cur_uid_lbl.setText("")
        self._reimport_btn.hide()
        self._set_state("empty")


# ─── CustomSendPanel ──────────────────────────────────────────────────────────
class CustomSendPanel(QWidget):
    """自定义消息发送面板"""

    def __init__(self, mqtt_mgr: "MQTTManager", parent=None):
        super().__init__(parent)
        self._mqtt = mqtt_mgr
        self._setup_ui()
        self._mqtt.sig_connected.connect(lambda: self._send_btn.setEnabled(True))
        self._mqtt.sig_disconnected.connect(lambda: self._send_btn.setEnabled(False))

    def _setup_ui(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 8, 0, 0)
        outer.setSpacing(0)

        grp = QGroupBox("📤  自定义消息")
        grp_layout = QVBoxLayout(grp)
        grp_layout.setSpacing(8)

        topic_row = QHBoxLayout()
        topic_lbl = QLabel("主题:")
        topic_lbl.setFixedWidth(42)
        self._topic_edit = QLineEdit()
        self._topic_edit.setPlaceholderText("输入发布主题")
        topic_row.addWidget(topic_lbl)
        topic_row.addWidget(self._topic_edit)

        self._msg_edit = QTextEdit()
        self._msg_edit.setPlaceholderText("输入消息内容（支持 JSON）...")
        self._msg_edit.setMaximumHeight(100)

        btn_row = QHBoxLayout()
        btn_row.addStretch()
        self._send_btn = _make_btn("⚡  发送", "#3EC78C", height=34)
        self._send_btn.setEnabled(False)
        self._send_btn.clicked.connect(self._do_send)
        btn_row.addWidget(self._send_btn)

        grp_layout.addLayout(topic_row)
        grp_layout.addWidget(self._msg_edit)
        grp_layout.addLayout(btn_row)

        outer.addWidget(grp)
        outer.addStretch()

    def _do_send(self):
        topic = self._topic_edit.text().strip()
        msg   = self._msg_edit.toPlainText().strip()
        if not topic:
            self._mqtt.sig_log.emit(MQTTManager._fmt("[WARN]", "请输入发布主题", "orange"))
            return
        if not msg:
            self._mqtt.sig_log.emit(MQTTManager._fmt("[WARN]", "消息内容不能为空", "orange"))
            return
        self._mqtt.publish(topic, msg)


# ─── OTAPanel ─────────────────────────────────────────────────────────────────
_JSON_STYLE_OK  = """QTextEdit {
    background-color: #1C2E26; color: #A8E6C3;
    font-family: Consolas, "Courier New", monospace; font-size: 11px;
    border-radius: 6px; border: 1px solid #2A4A3A; padding: 4px; }"""
_JSON_STYLE_ERR = """QTextEdit {
    background-color: #2E1C1C; color: #F4AFAF;
    font-family: Consolas, "Courier New", monospace; font-size: 11px;
    border-radius: 6px; border: 1px solid #5A2A2A; padding: 4px; }"""


class OTAPanel(QWidget):
    """跨平台 OTA 升级面板：填参数 → 预览 JSON → 单设备发送 / Excel 批量"""

    def __init__(self, mqtt_mgr: "MQTTManager", parent=None):
        super().__init__(parent)
        self._mqtt = mqtt_mgr
        self._confirmed_payload: str | None = None
        self._setup_ui()
        self._mqtt.sig_connected.connect(self._refresh_send_btn)
        self._mqtt.sig_disconnected.connect(self._refresh_send_btn)

    def _setup_ui(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 8, 0, 0)
        outer.setSpacing(8)

        # ── 参数输入区 ────────────────────────────────────────────
        param_grp = QGroupBox("🔧  OTA 参数配置")
        pg_layout = QVBoxLayout(param_grp)
        pg_layout.setSpacing(6)

        form = QFormLayout()
        form.setSpacing(6)
        form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        form.setFieldGrowthPolicy(QFormLayout.FieldGrowthPolicy.ExpandingFieldsGrow)

        self._uid_edit  = QLineEdit()
        self._uid_edit.setPlaceholderText("目标设备 UID（单设备发送时填写）")
        self._url_edit  = QLineEdit()
        self._url_edit.setPlaceholderText("http://...  固件下载地址")
        self._ver_edit  = QLineEdit("1.1")
        self._size_edit = QLineEdit()
        self._size_edit.setPlaceholderText("固件字节数（整数），如 170000")
        self._sha_edit  = QLineEdit()
        self._sha_edit.setPlaceholderText("64 位十六进制 SHA256 校验值")

        form.addRow("UID:",           self._uid_edit)
        form.addRow("espfileUrl:",    self._url_edit)
        form.addRow("espver:",        self._ver_edit)
        form.addRow("espfilesize:",   self._size_edit)
        form.addRow("espfilesha256:", self._sha_edit)

        # 参数变动 → 清除已确认 payload（防止旧数据发送）
        for edit in (self._url_edit, self._ver_edit, self._size_edit, self._sha_edit):
            edit.textChanged.connect(self._invalidate_preview)

        preview_row = QHBoxLayout()
        preview_row.addStretch()
        self._preview_btn = _make_btn("🔍  确认 / 预览 JSON", "#5B9BD5", height=32)
        self._preview_btn.setMinimumWidth(170)
        self._preview_btn.clicked.connect(self._do_preview)
        preview_row.addWidget(self._preview_btn)

        pg_layout.addLayout(form)
        pg_layout.addLayout(preview_row)
        outer.addWidget(param_grp)

        # ── JSON 预览 + 发送区 ────────────────────────────────────
        send_grp = QGroupBox("📋  发送预览")
        sg_layout = QVBoxLayout(send_grp)
        sg_layout.setContentsMargins(8, 6, 8, 8)
        sg_layout.setSpacing(6)

        self._topic_lbl = QLabel("发送主题: /func/4/{UID}/get")
        self._topic_lbl.setStyleSheet(
            "font-size: 11px; color: #5B9BD5; font-family: Consolas;")
        self._uid_edit.textChanged.connect(self._update_topic_lbl)

        self._json_preview = QTextEdit()
        self._json_preview.setReadOnly(True)
        self._json_preview.setMinimumHeight(100)
        self._json_preview.setMaximumHeight(130)
        self._json_preview.setPlaceholderText(
            "填写上方参数后，点击「确认 / 预览 JSON」查看即将发送的内容...")
        self._json_preview.setStyleSheet(_JSON_STYLE_OK)

        send_btn_row = QHBoxLayout()
        send_btn_row.addWidget(self._topic_lbl, 1)
        self._send_single_btn = _make_btn("🚀  发送至单设备", "#3EC78C", height=32)
        self._send_single_btn.setEnabled(False)
        self._send_single_btn.clicked.connect(self._do_send_single)
        send_btn_row.addWidget(self._send_single_btn)

        sg_layout.addWidget(self._json_preview)
        sg_layout.addLayout(send_btn_row)
        outer.addWidget(send_grp)

        # ── Excel 批量执行面板 ────────────────────────────────────
        self._batch_panel = BatchExecutionPanel(
            self._mqtt, self._get_payload, self)
        outer.addWidget(self._batch_panel)

    # ── 辅助 ──────────────────────────────────────────────────────────
    def _update_topic_lbl(self, uid: str):
        uid_part = uid.strip() if uid.strip() else "{UID}"
        self._topic_lbl.setText(f"发送主题: /func/4/{uid_part}/get")

    def _invalidate_preview(self):
        if self._confirmed_payload is not None:
            self._confirmed_payload = None
            self._send_single_btn.setEnabled(False)
            self._json_preview.setStyleSheet(_JSON_STYLE_OK)
            self._json_preview.clear()

    def _refresh_send_btn(self):
        self._send_single_btn.setEnabled(
            self._confirmed_payload is not None and self._mqtt.is_connected())

    def _get_payload(self) -> str | None:
        return self._confirmed_payload

    # ── 事件处理 ──────────────────────────────────────────────────────
    def _do_preview(self):
        url    = self._url_edit.text().strip()
        ver    = self._ver_edit.text().strip()
        size_s = self._size_edit.text().strip()
        sha    = self._sha_edit.text().strip()

        errors = []
        if not url:  errors.append("espfileUrl")
        if not ver:  errors.append("espver")
        if not sha:  errors.append("espfilesha256")

        size = 0
        if not size_s:
            errors.append("espfilesize")
        else:
            try:
                size = int(size_s)
                if size <= 0:
                    errors.append("espfilesize（需为正整数）")
            except ValueError:
                errors.append("espfilesize（需为整数）")

        if errors:
            self._json_preview.setStyleSheet(_JSON_STYLE_ERR)
            self._json_preview.setPlainText(
                "⚠️ 以下字段缺失或格式错误:\n  " + "\n  ".join(errors))
            self._confirmed_payload = None
            self._send_single_btn.setEnabled(False)
            return

        self._confirmed_payload = _build_ota_payload(url, ver, size, sha, pretty=False)
        pretty_json             = _build_ota_payload(url, ver, size, sha, pretty=True)
        self._json_preview.setStyleSheet(_JSON_STYLE_OK)
        self._json_preview.setPlainText(pretty_json)
        self._send_single_btn.setEnabled(self._mqtt.is_connected())

    def _do_send_single(self):
        if self._confirmed_payload is None:
            return
        uid = self._uid_edit.text().strip()
        if not uid:
            self._json_preview.setStyleSheet(_JSON_STYLE_ERR)
            self._json_preview.setPlainText("⚠️ 请先填写 UID 字段（单设备发送专用）")
            return
        topic = f"/func/4/{uid}/get"
        self._mqtt.publish(topic, self._confirmed_payload)


# ─── MQTTWidget ───────────────────────────────────────────────────────────────
class MQTTWidget(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("featurePage")
        self._mqtt = MQTTManager(self)
        self._mqtt.sig_connected.connect(self._on_connected)
        self._mqtt.sig_disconnected.connect(self._on_disconnected)
        self._mqtt.sig_log.connect(self._append_log)
        self._connected = False
        self._setup_ui()

    # ── UI ────────────────────────────────────────────────────────────
    def _setup_ui(self):
        main = QVBoxLayout(self)
        main.setContentsMargins(18, 10, 18, 14)
        main.setSpacing(10)

        # 标题行
        title_row = QHBoxLayout()
        icon_lbl = QLabel("📡")
        icon_lbl.setStyleSheet("font-size: 26px;")
        title_lbl = QLabel("高速转发  —  MQTT 客户端")
        title_lbl.setObjectName("featureTitle")
        title_row.addWidget(icon_lbl)
        title_row.addWidget(title_lbl)
        title_row.addStretch()
        main.addLayout(title_row)

        # 连接配置
        conn_grp = QGroupBox("⚙️  连接配置")
        conn_layout = QHBoxLayout(conn_grp)
        conn_layout.setSpacing(14)

        left_form = QFormLayout()
        left_form.setSpacing(6)
        left_form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        self._host_edit = QLineEdit("broker.emqx.io")
        self._host_edit.setPlaceholderText("MQTT Broker 地址")
        self._port_spin = QSpinBox()
        self._port_spin.setRange(1, 65535)
        self._port_spin.setValue(1883)
        self._port_spin.setFixedWidth(90)
        self._cid_edit = QLineEdit()
        self._cid_edit.setPlaceholderText("留空则自动生成")
        left_form.addRow("Broker:", self._host_edit)
        left_form.addRow("端口:", self._port_spin)
        left_form.addRow("Client ID:", self._cid_edit)

        right_form = QFormLayout()
        right_form.setSpacing(6)
        right_form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        self._user_edit = QLineEdit()
        self._user_edit.setPlaceholderText("用户名（可选）")
        self._pass_edit = QLineEdit()
        self._pass_edit.setPlaceholderText("密码（可选）")
        self._pass_edit.setEchoMode(QLineEdit.EchoMode.Password)
        self._ka_spin = QSpinBox()
        self._ka_spin.setRange(1, 3600)
        self._ka_spin.setValue(60)
        self._ka_spin.setFixedWidth(90)
        right_form.addRow("用户名:", self._user_edit)
        right_form.addRow("密码:", self._pass_edit)
        right_form.addRow("心跳(s):", self._ka_spin)

        ctrl_col = QVBoxLayout()
        ctrl_col.setSpacing(10)
        status_row = QHBoxLayout()
        self._dot_lbl = QLabel("⚫")
        self._dot_lbl.setStyleSheet("font-size: 16px;")
        self._status_lbl = QLabel("未连接")
        self._status_lbl.setStyleSheet("font-size: 13px; color: #7CB89C;")
        status_row.addWidget(self._dot_lbl)
        status_row.addWidget(self._status_lbl)
        status_row.addStretch()
        self._conn_btn = QPushButton("🔌  连接")
        self._conn_btn.setObjectName("primaryButton")
        self._conn_btn.setFixedWidth(110)
        self._conn_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self._conn_btn.clicked.connect(self._toggle_connection)
        ctrl_col.addLayout(status_row)
        ctrl_col.addWidget(self._conn_btn)
        ctrl_col.addStretch()

        conn_layout.addLayout(left_form, 4)
        conn_layout.addLayout(right_form, 4)
        conn_layout.addLayout(ctrl_col, 3)
        main.addWidget(conn_grp)

        # Tab 切换按钮
        tab_bar = QHBoxLayout()
        tab_bar.setSpacing(0)
        self._tab_custom_btn = QPushButton("📤  自定义发送")
        self._tab_ota_btn    = QPushButton("🛠️  跨平台 OTA")
        for btn in (self._tab_custom_btn, self._tab_ota_btn):
            btn.setFixedHeight(34)
            btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self._tab_custom_btn.clicked.connect(lambda: self._switch_tab(0))
        self._tab_ota_btn.clicked.connect(lambda: self._switch_tab(1))
        tab_bar.addWidget(self._tab_custom_btn)
        tab_bar.addWidget(self._tab_ota_btn)
        tab_bar.addStretch()
        main.addLayout(tab_bar)

        # 堆叠内容区
        self._stack = QStackedWidget()

        # Page 0：自定义发送
        self._custom_panel = CustomSendPanel(self._mqtt, self._stack)
        self._stack.addWidget(self._custom_panel)

        # Page 1：跨平台 OTA（带纵向滚动）
        ota_scroll = QScrollArea()
        ota_scroll.setWidgetResizable(True)
        ota_scroll.setFrameShape(QFrame.Shape.NoFrame)
        ota_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        ota_scroll.setStyleSheet("QScrollArea { background: transparent; }")
        self._ota_panel = OTAPanel(self._mqtt, ota_scroll)
        ota_scroll.setWidget(self._ota_panel)
        self._stack.addWidget(ota_scroll)

        main.addWidget(self._stack, 1)
        self._switch_tab(0)

        # 日志区
        log_grp = QGroupBox("🖥️  消息日志")
        log_layout = QVBoxLayout(log_grp)
        log_layout.setContentsMargins(6, 4, 6, 6)
        log_layout.setSpacing(4)

        log_header = QHBoxLayout()
        log_header.addStretch()
        clr_btn = QPushButton("清空")
        clr_btn.setFixedWidth(58)
        clr_btn.setStyleSheet("""
            QPushButton {
                background: #274035; color: #5EF09E;
                border: 1px solid #3A6050; border-radius: 5px;
                padding: 2px 6px; font-size: 11px;
            }
            QPushButton:hover { background: #3A5045; }
        """)
        clr_btn.clicked.connect(self._clear_log)
        log_header.addWidget(clr_btn)

        self._log_edit = QTextEdit()
        self._log_edit.setObjectName("logTextEdit")
        self._log_edit.setReadOnly(True)
        self._log_edit.setMinimumHeight(100)
        self._log_edit.setSizePolicy(
            QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        log_layout.addLayout(log_header)
        log_layout.addWidget(self._log_edit)
        main.addWidget(log_grp)

    # ── Tab 切换 ──────────────────────────────────────────────────────
    _TAB_ACTIVE = """QPushButton {
        background-color: #3EC78C; color: white; border: none;
        border-radius: 8px 8px 0 0; font-size: 13px; font-weight: bold;
        padding: 6px 22px; border-bottom: 3px solid #24A56B; }"""
    _TAB_INACTIVE = """QPushButton {
        background-color: #E8F5EE; color: #5A9B75;
        border: 1px solid #C8DDD4; border-bottom: none;
        border-radius: 8px 8px 0 0; font-size: 13px; font-weight: bold;
        padding: 6px 22px; }
    QPushButton:hover { background-color: #D4EDE0; }"""

    def _switch_tab(self, idx: int):
        self._stack.setCurrentIndex(idx)
        self._tab_custom_btn.setStyleSheet(
            self._TAB_ACTIVE if idx == 0 else self._TAB_INACTIVE)
        self._tab_ota_btn.setStyleSheet(
            self._TAB_ACTIVE if idx == 1 else self._TAB_INACTIVE)

    # ── 连接控制 ──────────────────────────────────────────────────────
    def _toggle_connection(self):
        if self._connected:
            self._mqtt.disconnect_broker()
        else:
            host = self._host_edit.text().strip()
            if not host:
                self._append_log(MQTTManager._fmt("[WARN]", "请填写 Broker 地址", "orange"))
                return
            self._conn_btn.setEnabled(False)
            self._conn_btn.setText("连接中...")
            self._mqtt.connect_broker(
                host=host,
                port=self._port_spin.value(),
                client_id=self._cid_edit.text().strip(),
                username=self._user_edit.text().strip(),
                password=self._pass_edit.text(),
                keepalive=self._ka_spin.value(),
            )

    def _on_connected(self):
        self._connected = True
        self._dot_lbl.setText("🟢")
        self._status_lbl.setText("已连接")
        self._conn_btn.setText("🔌  断开")
        self._conn_btn.setEnabled(True)
        self._conn_btn.setStyleSheet("""
            QPushButton {
                background-color: #FF6B6B; color: white;
                border: none; border-radius: 10px;
                padding: 10px 20px; font-size: 13px; font-weight: bold;
            }
            QPushButton:hover { background-color: #E85555; }
        """)

    def _on_disconnected(self):
        self._connected = False
        self._dot_lbl.setText("⚫")
        self._status_lbl.setText("未连接")
        self._conn_btn.setText("🔌  连接")
        self._conn_btn.setEnabled(True)
        self._conn_btn.setStyleSheet("")

    def _append_log(self, html_line: str):
        self._log_edit.append(html_line)
        sb = self._log_edit.verticalScrollBar()
        sb.setValue(sb.maximum())

    def _clear_log(self):
        self._log_edit.clear()

